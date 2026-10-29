"""Ticket export formatting and authorization regression tests."""
import json
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.core.exceptions import BusinessError
from app.libs.ticket_export_xlsx import TicketWorkbook, format_value
from app.schemas.permission import EffectivePrincipal
from app.services.ticket_export_service import TicketExportService, signature


def field(kind="single_line_text", **kwargs):
    return dict(id=None, key="ticket_number", source="system", slot_column=None,
                name="Number", field_type=kind, **kwargs)


def principal(scope="self", **kwargs):
    return EffectivePrincipal(user_id=7, tenant_id=1, data_scopes={"ticket": scope},
        permissions=list(TicketExportService.permissions), **kwargs)


class Session:
    active = False

    async def __aenter__(self):
        self.active = True
        return self

    async def __aexit__(self, *args):
        self.active = False


def test_formula_text_and_zero_are_preserved():
    book = TicketWorkbook([field(), field("number")], "en", "UTC")
    book.append(['=HYPERLINK("https://example.org")', 0])
    book.append(["0001234567890123456789", Decimal("1234567890123456.78")])
    sheet = load_workbook(BytesIO(book.finish())).active
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 0
    assert sheet["A3"].value == "0001234567890123456789"
    assert sheet["B3"].data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B3"


def test_empty_file_contains_headers_and_timezone():
    book = TicketWorkbook([field("datetime")], "zh", "Asia/Shanghai")
    sheet = load_workbook(BytesIO(book.finish())).active
    assert sheet.title == "工单"
    assert sheet["A1"].value == "Number（Asia/Shanghai）"
    assert sheet.max_row == 1


def test_dates_options_html_files():
    assert format_value(field("datetime"), datetime(2026, 9, 13, tzinfo=timezone.utc), "en", "Asia/Shanghai") == "2026-09-13 08:00:00"
    assert format_value(field("date"), "2026-09-13", "en", "America/New_York") == "2026-09-13"
    assert format_value(field("time"), "12:15:03.123", "en", "UTC") == "12:15:03.123"
    options = [{"value": "b", "label": "Bee", "sort_order": 1}, {"value": "a", "label": "Aye", "sort_order": 0}]
    assert format_value(field("multi_select", options=options), ["b", "a", "missing"], "en", "UTC") == "Aye\nBee\nUnavailable"
    assert format_value(field("rich_text"), '<p>Hello<br>world</p><script>secret()</script><img src="secret">', "en", "UTC") == "Hello\nworld"
    assert format_value(field("file"), [{"name": "report.pdf", "url": "secret"}], "en", "UTC") == "report.pdf"


def test_tree_full_paths():
    nodes = [{"id": 1, "parent_id": None, "label": "Parent", "value": "p"},
             {"id": 2, "parent_id": 1, "label": "Child", "value": "c"}]
    assert format_value(field("single_select_tree", tree_nodes=nodes), "c", "en", "UTC") == "Parent / Child"


def test_rejects_limits_without_truncating():
    book = TicketWorkbook([field()], "en", "UTC")
    with pytest.raises(BusinessError) as error:
        book.cell("a" * 32768, 0)
    assert error.value.code == "EXPORT_CELL_LIMIT_0"
    assert len(book.cell("a" * 32767, 0).value) == 32767
    book.count = 50000
    with pytest.raises(BusinessError) as error:
        book.append(["next"])
    assert error.value.code == "EXPORT_ROW_LIMIT"
    book.finish()
    with pytest.raises(BusinessError):
        TicketWorkbook([field()] * 201, "en", "UTC")


def test_field_identity_and_type_invalidation():
    frozen = field()
    TicketExportService.validate_fields([frozen], [dict(frozen, name="Renamed")])
    for current in ([], [dict(frozen, field_type="number")], [dict(frozen, slot_column="str_1")]):
        with pytest.raises(BusinessError):
            TicketExportService.validate_fields([frozen], current)


@pytest.mark.asyncio
async def test_signed_configuration_cannot_be_tampered_or_shared():
    prepared = {"tenant_id": 1, "employee_id": 7}
    signed = signature(json.dumps(prepared))
    with pytest.raises(BusinessError):
        await TicketExportService().freeze(principal(), dict(prepared=json.dumps(dict(prepared, tenant_id=2)), signature=signed, fields=["key:ticket_number"]))
    other = dict(prepared, employee_id=9)
    with pytest.raises(BusinessError):
        await TicketExportService().freeze(principal(), dict(prepared=json.dumps(other), signature=signature(json.dumps(other)), fields=["key:ticket_number"]))


@pytest.mark.asyncio
async def test_scope_intersection_does_not_expand():
    frozen = {"access": principal("group", group_ids=[2]).model_dump(), "peer_ids": [8]}
    with patch("app.services.ticket_export_service.DataScopeService.get_group_peer_employee_ids", AsyncMock(return_value=[99])):
        predicate = await TicketExportService.scope(None, principal("all"), frozen)
    sql = str(predicate.compile(dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True}))
    assert "99" not in sql
    assert "8" in sql and "2" in sql


@pytest.mark.asyncio
async def test_download_checks_every_record_and_empty_fields():
    service = TicketExportService()
    session = Session()
    config = {"language": "en", "fields": [field()], "access": principal().model_dump(), "peer_ids": []}
    with patch("app.services.ticket_export_service.AsyncSessionLocal", return_value=session), \
         patch.object(service, "catalog", AsyncMock(return_value=[field()])), \
         patch.object(service, "scope", AsyncMock(return_value=None)), \
         patch("app.services.ticket_export_service.Repo.rows", AsyncMock(return_value=[1])):
        assert not await service.can_download(principal(), {"config": config}, {"ticket_ids": [1, 2]})
        assert await service.can_download(principal(), {"config": config}, {"ticket_ids": []})
    assert not session.active


@pytest.mark.asyncio
async def test_generation_fixed_ids_and_closed_session_at_checkpoints():
    service = TicketExportService()
    session = Session()
    config = dict(language="en", fields=[field()], dependencies=[], query={},
                  submitted_at="2026-09-13T00:00:00+00:00", suffix="12345678")
    async def checkpoint():
        assert not session.active
    with patch("app.services.ticket_export_service.AsyncSessionLocal", return_value=session), \
         patch.object(service, "catalog", AsyncMock(return_value=[field()])), \
         patch.object(service, "scope", AsyncMock(return_value=None)), \
         patch.object(service, "row_values", AsyncMock(return_value=[["001"]])), \
         patch("app.services.ticket_export_service.ExportTaskService.current", AsyncMock(return_value=principal())), \
         patch("app.services.ticket_export_service.TicketRepository.query_paginated", AsyncMock(return_value=([1, 2], 0))) as query, \
         patch("app.services.ticket_export_service.Repo.rows", AsyncMock(return_value=[SimpleNamespace(id=1)])) as read:
        result = await service.generate({"config": config, "language": "en", "timezone": "UTC"}, principal(), checkpoint)
    assert result.manifest == {"ticket_ids": [1]}
    assert result.row_count == 1 and result.scope_reduced
    assert read.call_args.args[2] == [1, 2]
    assert query.call_args.kwargs["include_count"] is False
    assert query.call_args.kwargs["per_page"] == 50001
    assert load_workbook(BytesIO(result.data)).active["A2"].value == "001"


def test_aborted_workbook_cleans_spooled_data():
    from pathlib import Path
    book = TicketWorkbook([field()], 'en', 'UTC')
    path = Path(book.sheet._writer.out)
    assert path.exists()
    book.close()
    assert not path.exists()


def test_file_size_limit_fails_the_whole_workbook(monkeypatch):
    monkeypatch.setattr('app.libs.ticket_export_xlsx.MAX_BYTES', 1)
    book = TicketWorkbook([field()], 'en', 'UTC')
    with pytest.raises(BusinessError) as error:
        book.finish()
    assert error.value.code == 'EXPORT_FILE_LIMIT'
    book.close()


def test_unknown_filters_are_rejected_instead_of_broadening_query():
    with pytest.raises(BusinessError):
        TicketExportService.validate_query({'temp_conditions': [{'field_key': 'deleted', 'operator': 'eq', 'value': 'secret'}]}, [field()])
    with pytest.raises(BusinessError):
        TicketExportService.validate_query({'temp_conditions': [{'field_key': 'ticket_number', 'operator': 'unsupported'}]}, [field()])


@pytest.mark.asyncio
async def test_submission_rejects_known_row_limit():
    service = TicketExportService()
    prepared = dict(tenant_id=1, employee_id=7, fields=[field()], dependencies=[], language='en', query={}, summary={'en': 'Tickets'})
    with patch('app.services.ticket_export_service.AsyncSessionLocal', return_value=Session()), \
         patch.object(service, 'catalog', AsyncMock(return_value=[field()])), \
         patch.object(service, 'scope', AsyncMock(return_value=None)), \
         patch('app.services.ticket_export_service.TicketRepository.query_paginated', AsyncMock(return_value=(list(range(50001)), 0))) as query:
        with pytest.raises(BusinessError) as error:
            await service.freeze(principal(), dict(prepared=json.dumps(prepared), signature=signature(json.dumps(prepared)), fields=['key:ticket_number']))
    assert error.value.code == 'EXPORT_ROW_LIMIT'
    assert query.call_args.kwargs['include_count'] is False
    assert query.call_args.kwargs['per_page'] == 50001


@pytest.mark.asyncio
async def test_export_config_route_rejects_view_only_principal():
    from httpx import ASGITransport, AsyncClient
    from app.main import _fastapi_app
    from app.db.deps import get_current_principal_short_session
    viewer = EffectivePrincipal(user_id=7, tenant_id=1, permissions=['ticket.workspace.view'])
    previous = _fastapi_app.dependency_overrides.copy()
    _fastapi_app.dependency_overrides[get_current_principal_short_session] = lambda: viewer
    try:
        async with AsyncClient(transport=ASGITransport(app=_fastapi_app), base_url='http://test') as client:
            response = await client.post('/api/v1/tickets/export-config', json={'query': {}})
        assert response.status_code == 403
    finally:
        _fastapi_app.dependency_overrides = previous


@pytest.mark.asyncio
async def test_catalog_includes_shared_fields_and_disambiguates_headers():
    items = [dict(field(), domain='ticket', name='Same', status='active', sort_order=1),
             dict(field(), id=8, key='shared', domain='shared_pool', source='custom', applicable_modules=['ticket'], slot_column='str_1', name='Same', status='active', sort_order=2),
             dict(field(), id=9, key='other', domain='ticket', source='custom', slot_column='str_2', name='Same（2）', status='active', sort_order=3),
             dict(field(), id=10, key='off', domain='ticket', source='custom', name='Hidden', status='inactive')]
    with patch('app.services.ticket_export_service.FdFieldDefinitionService.get_unified_list', AsyncMock(return_value={'items': items})):
        result = await TicketExportService.catalog(None, 1, 'en')
    assert len(result) == 3
    assert len({item['name'] for item in result}) == 3
    assert result[1]['domain'] == 'shared_pool'


@pytest.mark.asyncio
async def test_preparation_freezes_view_and_null_group_without_pagination():
    from app.schemas.ticket_export import TicketExportPrepare
    service = TicketExportService()
    fields = [field(), dict(field(), id=7, key='custom', source='custom', slot_column='str_1')]
    view = SimpleNamespace(name='Saved view', conditions=[dict(field_key='ticket_number', operator='contains', value='TK')], condition_logic='or', group_field_id=7)
    with patch('app.services.ticket_export_service.AsyncSessionLocal', return_value=Session()), \
         patch.object(TicketExportService, 'catalog', AsyncMock(return_value=fields)), \
         patch.object(TicketExportService, 'scope', AsyncMock(return_value=None)), \
         patch('app.services.ticket_export_service.Repo.view', AsyncMock(return_value=view)), \
         patch('app.services.ticket_export_service.TicketRepository.query_paginated', AsyncMock(return_value=([], 30))) as query:
        result = await service.prepare(principal(), TicketExportPrepare(query={'view_id': 1, 'page': 3, 'group_value': '__EMPTY__'}))
    frozen = json.loads(result['prepared'])['query']
    assert frozen['group_value'] == '__EMPTY__'
    assert frozen['group_field_col'] == 'str_1'
    assert frozen['view_conditions'][0]['value'] == 'TK'
    assert 'page' not in frozen and 'view_id' not in frozen
    assert query.call_args.kwargs['per_page'] == 1
    assert result['signature'] == signature(result['prepared'])


@pytest.mark.asyncio
async def test_frozen_json_preserves_numbers_and_option_labels():
    service = TicketExportService()
    selected = dict(field(), type_config={'min': 1.0, 'max': 9007199254740993})
    prepared = json.dumps(dict(tenant_id=1, employee_id=7, fields=[selected], dependencies=[],
        language='en', query={}, summary={'en': 'Tickets'}))
    payload = json.loads(json.dumps(dict(prepared=prepared, signature=signature(prepared), fields=['key:ticket_number'])))
    with patch('app.services.ticket_export_service.AsyncSessionLocal', return_value=Session()), \
         patch.object(service, 'catalog', AsyncMock(return_value=[dict(selected, name='Renamed')])), \
         patch.object(service, 'scope', AsyncMock(return_value=None)), \
         patch('app.services.ticket_export_service.DataScopeService.get_group_peer_employee_ids', AsyncMock(return_value=[8])), \
         patch('app.services.ticket_export_service.TicketRepository.query_paginated', AsyncMock(return_value=([], 50000))):
        frozen = await service.freeze(principal(), payload)
    assert frozen.config['fields'][0]['name'] == 'Number'
    assert frozen.config['fields'][0]['type_config']['max'] == 9007199254740993
    assert frozen.config['peer_ids'] == [8]
